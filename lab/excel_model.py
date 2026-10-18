"""Build outputs/growth_model.xlsx: segment LTV, payback, channel ROAS, and win-back ROI.

Blue cells are assumptions to edit. Black cells come from the Instacart data or are
formulas. Instacart has no prices or ad spend, so every rupee figure is an assumption.
"""

import json

import duckdb
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from lab.config import DB_PATH, OUTPUTS

ASSUMPTION = Font(color="1F4FD8", bold=True)
HEADER = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="23262F")
NOTE = Font(italic=True, color="6B7280")
INR = '"₹"#,##0'
INR2 = '"₹"#,##0.00'
PCT = "0.0%"


def _header(ws, row, labels):
    for col, label in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.font, cell.fill = HEADER, HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build() -> str:
    con = duckdb.connect(str(DB_PATH), read_only=True)
    segments = con.sql("select * from m_segments order by segment").fetchall()
    con.close()
    exp = json.loads((OUTPUTS / "experiments.json").read_text())["rehabit_test"]["design"]

    wb = Workbook()

    ws = wb.active
    ws.title = "Inputs"
    _widths(ws, [44, 16, 70])
    ws["A1"] = "Assumptions (blue cells are editable)"
    ws["A1"].font = Font(bold=True, size=13)
    inputs = [
        ("Average order value", 450, INR, "Assumption. Instacart has no prices."),
        ("Contribution per order", 30, INR, "Assumption. What one order earns after delivery and discounts."),
        ("Win-back incentive cost per redeemed order", 75, INR, "Assumption. Coupon value paid only when used."),
        ("Message cost per user (push or WhatsApp)", 0.5, INR2, "Assumption."),
        ("LTV horizon (months)", 12, "0", "Assumption."),
    ]
    for i, (label, value, fmt, note) in enumerate(inputs, start=3):
        ws.cell(row=i, column=1, value=label)
        c = ws.cell(row=i, column=2, value=value)
        c.font, c.number_format = ASSUMPTION, fmt
        ws.cell(row=i, column=3, value=note).font = NOTE
    aov, contrib, incentive, msg_cost, horizon = (f"Inputs!$B${r}" for r in range(3, 8))

    ws = wb.create_sheet("Segments")
    _widths(ws, [32, 12, 12, 14, 12, 16, 18, 18, 18, 14])
    ws["A1"] = "Segment value. Users, orders per 30 days, and lapse rate come from the data."
    ws["A1"].font = Font(bold=True)
    _header(ws, 3, ["Segment", "Users", "Orders per 30 days", "Lapse rate (30+ day gap)",
                    "Monthly churn", "Contribution per month", "LTV over horizon",
                    "Max CAC for 3x LTV/CAC", "Payback at ₹300 CAC (months)", "Share of orders"])
    for i, seg in enumerate(segments, start=4):
        name, users, _, share_orders, _, _, per30, lapse = seg[:8]
        churn_default = {"1.": 0.03, "2.": 0.06, "3.": 0.12, "4.": 0.25}[name[:2]]
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=int(users)).number_format = "#,##0"
        ws.cell(row=i, column=3, value=round(float(per30), 3)).number_format = "0.00"
        ws.cell(row=i, column=4, value=round(float(lapse), 4)).number_format = PCT
        c = ws.cell(row=i, column=5, value=churn_default)
        c.font, c.number_format = ASSUMPTION, PCT
        ws.cell(row=i, column=6, value=f"=C{i}*{contrib}").number_format = INR
        ws.cell(row=i, column=7, value=f"=F{i}*(1-(1-E{i})^{horizon})/E{i}").number_format = INR
        ws.cell(row=i, column=8, value=f"=G{i}/3").number_format = INR
        ws.cell(row=i, column=9, value=f"=300/F{i}").number_format = "0.0"
        ws.cell(row=i, column=10, value=round(float(share_orders), 4)).number_format = PCT
    last = 3 + len(segments)
    ws.cell(row=last + 2, column=1,
            value="Monthly churn is an assumption: Instacart only shows users up to their last order, "
                  "so true churn cannot be measured. Replace it with real churn when you have it.").font = NOTE

    ws = wb.create_sheet("Channels")
    _widths(ws, [24, 14, 12, 14, 12, 12, 16, 16, 14])
    ws["A1"] = "Acquisition channels. All figures are illustrative assumptions."
    ws["A1"].font = Font(bold=True)
    _header(ws, 3, ["Channel", "Spend", "Installs", "First-order rate", "Weekly-regular share",
                    "CAC (per first order)", "First-order ROAS", "12-month LTV/CAC", "Payback (months)"])
    channels = [
        ("Meta app installs", 500000, 25000, 0.18, 0.20),
        ("Google app campaigns", 500000, 20000, 0.24, 0.25),
        ("Referral (₹100 each side)", 300000, 3000, 0.60, 0.35),
        ("Influencer / content", 200000, 6000, 0.15, 0.15),
    ]
    for i, (name, spend, installs, rate, weekly_share) in enumerate(channels, start=4):
        ws.cell(row=i, column=1, value=name)
        for col, value, fmt in [(2, spend, INR), (3, installs, "#,##0"), (4, rate, PCT), (5, weekly_share, PCT)]:
            c = ws.cell(row=i, column=col, value=value)
            c.font, c.number_format = ASSUMPTION, fmt
        ws.cell(row=i, column=6, value=f"=B{i}/(C{i}*D{i})").number_format = INR
        ws.cell(row=i, column=7, value=f"={aov}/F{i}").number_format = "0.00x"
        blended = f"(E{i}*Segments!$G$4+(1-E{i})*AVERAGE(Segments!$G$5:$G${last}))"
        ws.cell(row=i, column=8, value=f"={blended}/F{i}").number_format = "0.00x"
        ws.cell(row=i, column=9, value=f"=F{i}/({blended}/{horizon})").number_format = "0.0"
    ws.cell(row=10, column=1,
            value="12-month LTV/CAC blends the weekly-regular LTV (by weekly-regular share) "
                  "with the average LTV of the other three segments.").font = NOTE

    ws = wb.create_sheet("Win-back")
    _widths(ws, [52, 16, 64])
    ws["A1"] = "Post-lapse re-habit campaign for weekly and fortnightly users"
    ws["A1"].font = Font(bold=True)
    rows = [
        ("Eligible users (from data)", exp["eligible_users"], "#,##0", False,
         "Weekly or fortnightly users who just came back after a 30+ day gap."),
        ("Baseline: next order within 30 days (from data)", round(exp["baseline_rate"], 4), PCT, False,
         "What happens with no campaign."),
        ("Lift from the test (absolute)", 0.03, PCT, True, "Use the A/B readout. 3 points is the simulated lift."),
        ("Extra orders per retained user over 3 months", 6, "0.0", True,
         "Assumption, roughly 3 months of fortnightly ordering."),
        ("Incremental users retained", "=B3*B5", "#,##0", False, ""),
        ("Incremental contribution", f"=B7*B6*{contrib}", INR, False, ""),
        ("Message cost", f"=B3*{msg_cost}", INR, False, ""),
        ("Incentive cost (all redeemers, incl. those who would return anyway)",
         f"=B3*(B4+B5)*{incentive}", INR, False, "The real cost of coupons: most redeemers would have ordered anyway."),
        ("Net return", "=B8-B9-B10", INR, False, ""),
        ("Break-even incentive per redeemed order", f"=(B8-B9)/(B3*(B4+B5))", INR, False,
         "If the coupon costs more than this, the campaign loses money."),
    ]
    for i, (label, value, fmt, editable, note) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=label)
        c = ws.cell(row=i, column=2, value=value)
        c.number_format = fmt
        if editable:
            c.font = ASSUMPTION
        ws.cell(row=i, column=3, value=note).font = NOTE

    path = OUTPUTS / "growth_model.xlsx"
    wb.save(path)
    return str(path)


if __name__ == "__main__":
    print(build())
